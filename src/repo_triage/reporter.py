"""Build a deterministic briefing shared by terminal and Markdown output."""

import html
import importlib
import io
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import TYPE_CHECKING

from rich.console import Console
from rich.markdown import Markdown

from repo_triage.llm_client import rank_triage
from repo_triage.models import Issue, IssueTriage

if TYPE_CHECKING:
    from docx.text.paragraph import Paragraph

_NOTICE = (
    "AI-assisted suggestions, not authoritative severity determinations. "
    "Maintainers should verify evidence and impact."
)


def _escape(text: str) -> str:
    """Keep untrusted text on one line and escape Markdown and HTML markup."""
    text = html.escape(" ".join(text.split()), quote=False)
    return re.sub(r"([\\`*_{}\[\]()#+.!|>-])", r"\\\1", text)


def _utc_timestamp(value: datetime) -> str:
    """Format an aware timestamp in UTC, rejecting ambiguous local times."""
    if value.utcoffset() is None:
        raise ValueError("Report timestamps must include a timezone.")
    return value.astimezone(timezone.utc).isoformat()


def build_digest(
    repository: str,
    issues: list[Issue],
    results: list[IssueTriage],
    requested_limit: int = 15,
    generated_at: datetime | None = None,
    since: datetime | None = None,
    limit_reason: str | None = None,
) -> str:
    """Build a Markdown briefing from validated triage without an LLM call.

    Args:
        repository: Validated owner/repo identifier.
        issues: Sampled source issues with original titles and URLs.
        results: Validated triage results matching the source issues exactly.
        requested_limit: Maximum issue count requested by the user.
        generated_at: Optional timezone-aware timestamp for reproducible output.
        since: Optional exclusive lower bound requested for issue updates.
        limit_reason: Whether the issue or page cap limited retrieval.

    Returns:
        A Markdown report with deterministic summary, priorities, and counts.

    Raises:
        ValueError: If source and result issue numbers do not match exactly.
    """
    sources = {issue.number: issue for issue in issues}
    numbers = [result.issue_number for result in results]
    if (
        len(sources) != len(issues)
        or len(numbers) != len(sources)
        or set(numbers) != set(sources)
    ):
        raise ValueError("Report requires one triage result per source issue.")
    ranked = rank_triage(results, issues)
    categories = Counter(result.category for result in ranked)
    urgencies = Counter(result.urgency for result in ranked)
    timestamp = generated_at or datetime.now(timezone.utc)
    requested_window = (
        f"updated_at > {_utc_timestamp(since)} (UTC; exclusive start, "
        "no explicit upper cutoff)."
        if since is not None
        else "none (no date filter)."
    )
    observed = [issue.updated_at for issue in issues]
    # Validate awareness before comparing timestamps or converting timezones.
    for updated_at in observed:
        _utc_timestamp(updated_at)
    observed_oldest = _utc_timestamp(min(observed)) if observed else "N/A"
    observed_newest = _utc_timestamp(max(observed)) if observed else "N/A"
    retrieval_status = {
        "issue_limit": "Issue limit reached; additional matching issues may "
        "exist. The requested window is not guaranteed to be fully covered.",
        "page_limit": "Two-page cap reached; additional matching issues may "
        "exist. The requested window is not guaranteed to be fully covered.",
    }.get(
        limit_reason,
        "No retrieval cap reached; coverage reflects the "
        "open issues returned by GitHub during this retrieval.",
    )
    lines = [
        f"# Monday briefing: {_escape(repository)}",
        "",
        "## Repository and analysis scope",
        "",
        f"- Repository: https://github.com/{repository}",
        f"- Analyzed: {len(ranked)} open issues (requested up to "
        f"{requested_limit}); pull requests excluded.",
        f"- Requested update window: {requested_window}",
        f"- Observed oldest update (analyzed issues, UTC): {observed_oldest}",
        f"- Observed newest update (analyzed issues, UTC): {observed_newest}",
        f"- Retrieval status: {retrieval_status}",
        "- Sorting: sampled by most recently updated first; priorities ranked "
        "by urgency, priority labels, comment count, then issue number.",
        "- Sampling limitation: at most two GitHub pages of 100 entries, "
        "with at most 15 actual issues analyzed. This is a recent-activity "
        "sample, not a full backlog audit; older urgent issues may be missed. "
        "Issue bodies are truncated to 800 characters.",
        f"- Generated at: {_utc_timestamp(timestamp)} (UTC).",
        "",
        "## Executive summary",
        "",
    ]
    if ranked:
        counts = ", ".join(
            f"{urgencies[level]} {level}"
            for level in ("critical", "high", "medium", "low")
        )
        lines.append(
            f"The {len(ranked)} sampled issues include "
            f"{categories['bug']} bug reports. Suggested urgency: {counts}. "
            f"Start by reviewing #{ranked[0].issue_number}."
        )
    else:
        lines.append(
            "No matching issues were returned within the sampled pages. "
            "No LLM analysis was needed; this does not establish that the "
            "repository has no open issues."
        )
    lines.extend(["", "## Top 5 priorities", ""])
    if not ranked:
        lines.append("No sampled issues to prioritize.")
    for index, result in enumerate(ranked[:5], start=1):
        source = sources[result.issue_number]
        # URLs come from the source, never from the LLM. Encode delimiters.
        url = str(source.html_url).replace("(", "%28").replace(")", "%29")
        lines.extend(
            [
                f"### {index}. #{source.number}: {_escape(source.title)}",
                "",
                f"- URL: {url}",
                f"- Last updated (UTC): {_utc_timestamp(source.updated_at)}",
                f"- Urgency: {result.urgency}",
                f"- Rationale: {_escape(result.rationale)}",
                f"- Next action: {_escape(result.next_action)}",
                "",
            ]
        )
    lines.extend(["", "## Category counts", ""])
    lines.extend(
        f"- {category}: {categories[category]}"
        for category in ("bug", "feature", "docs", "question", "other")
    )
    lines.extend(["", f"**{_NOTICE}**", ""])
    return "\n".join(lines)


def print_digest(digest: str, console: Console) -> None:
    """Render the shared Markdown briefing as concise Rich terminal output."""
    console.print(Markdown(digest))


class ExportError(Exception):
    """A safe report-export failure without raw library or file details."""


def check_export_support(output_format: str) -> None:
    """Validate the format and optional dependency before paid analysis."""
    if output_format not in {"md", "docx", "both"}:
        raise ExportError("Report format must be md, docx, or both.")
    if output_format in {"docx", "both"}:
        try:
            importlib.import_module("docx")
        except ImportError:
            raise ExportError(
                "Word export requires python-docx. Install it with: "
                "python -m pip install -e '.[word]'"
            ) from None


def _plain_text(text: str) -> str:
    """Decode only the escaping emitted by this report's Markdown builder."""
    if text.startswith("**") and text.endswith("**"):
        text = text[2:-2]
    return html.unescape(re.sub(r"\\([\\`*_{}\[\]()#+.!|>-])", r"\1", text))


def _add_link(paragraph: "Paragraph", url: str) -> None:
    """Add a clickable source URL without fetching or executing its content."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    link = OxmlElement("w:hyperlink")
    link.set(
        qn("r:id"),
        paragraph.part.relate_to(
            url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        ),
    )
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _docx_bytes(digest: str, generated_at: datetime) -> bytes:
    """Render the canonical report's headings, lists, text, and source links."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(12)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 14)):
        style = document.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 41, 55)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    bullets = document.styles["List Bullet"].paragraph_format
    bullets.left_indent, bullets.first_line_indent = Inches(0.5), Inches(-0.25)
    bullets.space_after = Pt(6)
    for line in digest.splitlines():
        if not line:
            continue
        heading = re.match(r"^(#{1,3}) (.*)$", line)
        if heading:
            level = len(heading[1]) - 1
            document.add_heading(_plain_text(heading[2]), level=level)
        else:
            bullet = line.startswith("- ")
            content = _plain_text(line[2:] if bullet else line)
            paragraph = document.add_paragraph(
                style="List Bullet" if bullet else "Normal"
            )
            link = re.fullmatch(r"(Repository|URL): (https?://\S+)", content)
            if link:
                paragraph.add_run(link[1] + ": ")
                _add_link(paragraph, link[2])
            else:
                paragraph.add_run(content)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = (
                bullet
                and content.startswith(
                    ("URL:", "Last updated (UTC):", "Urgency:", "Rationale:")
                )
            )
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    timestamp = generated_at.astimezone(timezone.utc).replace(tzinfo=None)
    document.core_properties.created = timestamp
    document.core_properties.modified = timestamp
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_digest(
    digest: str,
    repository: str,
    generated_at: datetime,
    output_format: str = "md",
) -> list[Path]:
    """Export one canonical report without overwriting existing files.

    Args:
        digest: Markdown produced from the validated triage results.
        repository: Repository identifier; only a sanitized basename is used.
        generated_at: Aware report-generation timestamp, also used in filenames.
        output_format: md, docx, or both.

    Returns:
        Paths written under reports/, sharing one basename for both formats.

    Raises:
        ExportError: On missing dependencies, rendering, or filesystem failure.
        ValueError: If the generation timestamp has no timezone.
    """
    check_export_support(output_format)
    _utc_timestamp(generated_at)
    timestamp = generated_at.astimezone(timezone.utc)
    name = re.sub(r"[^A-Za-z0-9_.-]+", "_", repository.rsplit("/", 1)[-1])
    name = name.strip("._-")[:100] or "repository"
    stem = f"{name}_triage_{timestamp:%Y-%m-%d}"
    payloads: dict[str, bytes] = {}
    if output_format in {"md", "both"}:
        payloads["md"] = digest.encode("utf-8")
    if output_format in {"docx", "both"}:
        try:
            payloads["docx"] = _docx_bytes(digest, timestamp)
        except Exception:
            raise ExportError("Could not create the Word report.") from None
    directory = Path("reports")
    try:
        directory.mkdir(parents=True, exist_ok=True)
        for attempt in range(100):
            suffix = "" if attempt == 0 else f"_{timestamp:%H%M%S}"
            if attempt > 1:
                suffix += f"_{attempt}"
            basename = stem + suffix
            if any(
                (directory / f"{basename}.{ext}").exists()
                for ext in ("md", "docx")
            ):
                continue
            created: list[Path] = []
            try:
                for extension, payload in payloads.items():
                    path = directory / f"{basename}.{extension}"
                    with path.open("xb") as stream:
                        created.append(path)
                        stream.write(payload)
            except Exception:
                for path in created:
                    path.unlink(missing_ok=True)
                raise
            return created
    except OSError:
        raise ExportError(
            "Could not write reports; check directory permissions and space."
        ) from None
    raise ExportError("Too many reports already exist for this timestamp.")
